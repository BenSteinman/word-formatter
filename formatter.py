from copy import deepcopy
from pathlib import Path
from docx import Document
from utils import backup_document
from utils import load_rules
from iterator import iter_blocks
from classifier import classify_paragraph, classify_table
from docx.text.paragraph import Paragraph
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.table import Table

def rules_apply(rules: dict) -> bool:
    """
    Checks whether any paragraph-level formatting rules are set.

    Args:
        rules (dict): Parsed rules from rules.toml.

    Returns:
        bool: True if at least one paragraph rule is active, False otherwise.
    """
    p = rules.get("paragraph", {})
    return any([
        p.get("font"),
        p.get("base_size"),
        p.get("all_bold"),
        p.get("no_bold"),
        p.get("all_italic"),
        p.get("no_italic"),
        p.get("all_underline"),
        p.get("no_underline"),
    ])

def apply_heading_rules(paragraph: Paragraph, level: int, rules: dict):
    """
    Applies heading-level formatting rules to a paragraph.

    Args:
        paragraph (Paragraph): The heading paragraph to format.
        level (int): The heading level (e.g. 1, 2).
        rules (dict): Parsed rules from rules.toml.
    """
    headings = rules.get("heading", {})
    h = headings.get(str(level), {})
    for run in paragraph.runs:
        if h.get("size"):   run.font.size = Pt(h["size"])
        if "bold" in h:     run.bold = h["bold"]
        if h.get("font"):   run.font.name = h["font"]
        if h.get("color"):
            hex_color = h["color"].lstrip("#")
            r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)
        if h.get("text_case") == "upper":
            run.text = run.text.upper()
        elif h.get("text_case") == "title":
            run.text = run.text.title()

def apply_run_rules(paragraph: Paragraph, rules: dict):
    """
    Applies paragraph-level run formatting rules (font, size, bold, italic, underline).

    Args:
        paragraph (Paragraph): The paragraph whose runs will be formatted.
        rules (dict): Parsed rules from rules.toml.
    """
    p = rules.get("paragraph", {})
    for run in paragraph.runs:
        if p.get("font"):        run.font.name = p["font"]
        if p.get("base_size"):   run.font.size = Pt(p["base_size"])
        if p.get("all_bold"):    run.bold = True
        if p.get("no_bold"):     run.bold = False
        if p.get("all_italic"):  run.italic = True
        if p.get("no_italic"):   run.italic = False
        if p.get("all_underline"): run.underline = True
        if p.get("no_underline"):  run.underline = False

def format_document(input_path: str):
    """
    Main entry point. Loads a Word document, applies all formatting rules
    from rules.toml, and saves the result to an Output/ folder next to the original.

    Args:
        input_path (str): Path to the input .docx file.
    """
    # TODO 1: Load the input document
    doc = Document(input_path)

    # TODO 2: Save a backup using backup_document() from utils.py
    backup_document(input_path)

    # TODO 3: Clone the source document to preserve styles, numbering, and theme, then clear the body
    new_doc = deepcopy(doc)
    sectPr = new_doc.element.body.find(qn('w:sectPr'))
    new_doc.element.body.clear()
    if sectPr is not None:
        new_doc.element.body.append(sectPr)

    # TODO 4: Iterate through blocks using iter_blocks()
    rules = load_rules()

    h1_count = 0
    h2_count = 0
    index = 0

    for block in iter_blocks(doc):
        if block["type"] == "paragraph":
            new_doc.element.body.append(deepcopy(block["object"]._element))
            new_paragraph = Paragraph(new_doc.element.body[-1], new_doc)
            meta = classify_paragraph(block["object"])

            apply_memo_header_rules(new_paragraph, rules, index)

            if meta["heading_level"] is not None:
                if rules.get("heading", {}).get("number_headings"):
                    if meta["heading_level"] == 1:
                        h1_count += 1
                        h2_count = 0
                        prefix = f"{h1_count}. "
                    elif meta["heading_level"] == 2:
                        h2_count += 1
                        prefix = f"{h1_count}.{h2_count}. "
                    else:
                        prefix = ""
                    if prefix and new_paragraph.runs:
                        new_paragraph.runs[0].text = prefix + new_paragraph.runs[0].text
                apply_heading_rules(new_paragraph, meta["heading_level"], rules)
            elif meta["is_list"]:
                pass  # preserve list formatting as-is for now
            else:
                if rules_apply(rules):
                    apply_run_rules(new_paragraph, rules)

        elif block["type"] == "table":
            new_doc.element.body.append(deepcopy(block["object"]._element))
            new_table = Table(new_doc.element.body[-1], new_doc)
            apply_table_rules(new_table, rules)

        index += 1

    # Document Wide Rules

    apply_footer_rules(new_doc, rules)
    apply_header_rules(new_doc, rules)
    apply_margin_rules(new_doc, rules)

    # TODO 5: Save the output document to an Output/ folder next to the original
    original = Path(input_path)
    output_dir = original.parent / "Output"
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / original.name

    # TODO 6: Print the path of the saved output document
    try:
        new_doc.save(str(output_path))
        print(f"Saved to: {output_path}")
    except PermissionError:
        print("Error: Please close the output file in Word before running the formatter.")

def apply_footer_rules(new_doc: Document, rules: dict):
    """
    Applies footer formatting rules including text, font, alignment, and page numbers.

    Args:
        new_doc (Document): The output document.
        rules (dict): Parsed rules from rules.toml.
    """
    f = rules.get("footer", {})
    section = new_doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]

    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    
    if f.get("text"):
        paragraph.text = f["text"]

    if paragraph.runs:
        if f.get("font"):       paragraph.runs[0].font.name = f["font"]
        if f.get("font_size"):  paragraph.runs[0].font.size = Pt(int(f["font_size"]))

    if f.get("text_alignment"):
        paragraph.alignment = alignments.get(f["text_alignment"].lower())

    if f.get("page_numbers"):
        # page numbers require raw XML field codes
        add_page_number(paragraph)

def add_page_number(paragraph: Paragraph):
    """
    Inserts a PAGE field code into a paragraph using raw XML.
    Used to add auto-updating page numbers to the footer.

    Args:
        paragraph (Paragraph): The paragraph to insert the page number into.
    """
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def apply_header_rules(new_doc: Document, rules: dict):
    """
    Applies header formatting rules including text, font, alignment, and logo image.

    Args:
        new_doc (Document): The output document.
        rules (dict): Parsed rules from rules.toml.
    """
    h = rules.get("header", {})
    section = new_doc.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]

    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }

    if h.get("text"):
        paragraph.text = h["text"]

    if paragraph.runs:
        if h.get("font"):      paragraph.runs[0].font.name = h["font"]
        if h.get("font_size"): paragraph.runs[0].font.size = Pt(int(h["font_size"]))

    if h.get("text_alignment"):
        paragraph.alignment = alignments.get(h["text_alignment"].lower())

    if h.get("logo"):
        run = paragraph.add_run()
        run.add_picture(h["logo"])

    if h.get("alignment"):
        paragraph.alignment = alignments.get(h["alignment"].lower())

def apply_margin_rules(doc: Document, rules: dict):
    """
    Applies margin and header distance rules to the document's first section.

    Args:
        doc (Document): The output document.
        rules (dict): Parsed rules from rules.toml.
    """
    m = rules.get("margins", {})
    section = doc.sections[0]
    if m.get("top"):             section.top_margin      = Inches(m["top"])
    if m.get("bottom"):          section.bottom_margin   = Inches(m["bottom"])
    if m.get("left"):            section.left_margin     = Inches(m["left"])
    if m.get("right"):           section.right_margin    = Inches(m["right"])
    if m.get("header_distance"): section.header_distance = Inches(m["header_distance"])

def apply_memo_header_rules(paragraph: Paragraph, rules: dict, index: int):
    """
    Detects and formats memo header paragraphs (To:, From:, Date:, Re:).

    Only checks the first 6 blocks. If a match is found, uppercases the label,
    preserves the value, and applies tab alignment and optional bold formatting.

    Args:
        paragraph (Paragraph): The paragraph to check and potentially format.
        rules (dict): Parsed rules from rules.toml.
        index (int): The block index in the document (0-based).
    """
    if index >= 6:
        return

    import re
    if not paragraph.runs or not re.match(r"^(To:|From:|Date:|Re:)", paragraph.runs[0].text):
        return

    m = rules.get("memo_header", {})

    full_text = paragraph.text
    if ":" not in full_text:
        return

    label, value = full_text.split(":", 1)
    label = label.strip().upper() + ":"
    value = value.strip()

    # clear existing runs and rewrite
    for run in paragraph.runs:
        run.text = ""
    paragraph.runs[0].text = label + "\t" + value
    if "bold" in m:
        paragraph.runs[0].bold = m["bold"]

    # set tab stop
    tab_stop_inch = m.get("tab_stop", 1.5)
    pPr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:pos"), str(int(tab_stop_inch * 1440)))
    tabs.append(tab)
    pPr.append(tabs)

def apply_table_rules(table: Table, rules: dict):
    """
    Applies table formatting rules including autofit, alternating row/column colors,
    header row color, and alignment.

    Args:
        table (Table): The table to format.
        rules (dict): Parsed rules from rules.toml.
    """
    t = rules.get("table", {})

    if t.get("autofit") == "contents":
        table.allow_autofit = True

        # remove table-level explicit width
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "auto")
        tblW.set(qn("w:w"), "0")

        # remove cell-level explicit widths
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.find(qn("w:tcPr"))
                if tcPr is not None:
                    tcW = tcPr.find(qn("w:tcW"))
                    if tcW is not None:
                        tcPr.remove(tcW)

    elif t.get("autofit") in ("window", "fixed"):
        table.allow_autofit = t["autofit"] == "window"

    primary = t.get("primary_color")
    secondary = t.get("secondary_color")

    if primary and secondary:
        if t.get("alternate") == "rows":
            for row_index, row in enumerate(table.rows):
                color = primary if row_index % 2 == 0 else secondary
                for cell in row.cells:
                    set_cell_color(cell, color)
        else:  # default to columns
            for col_index, column in enumerate(table.columns):
                color = primary if col_index % 2 == 0 else secondary
                for cell in column.cells:
                    set_cell_color(cell, color)

    if t.get("header_row_color"):
        for cell in table.rows[0].cells:
            set_cell_color(cell, t["header_row_color"])


    table_alignments = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
    }

    if t.get("alignment"):
        table.alignment = table_alignments.get(t["alignment"].lower())

def set_cell_color(cell, hex_color: str):
    """
    Sets the background fill color of a table cell using raw XML.

    Args:
        cell: A python-docx table cell object.
        hex_color (str): A hex color string (e.g. "#FF0000" or "FF0000").
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)                